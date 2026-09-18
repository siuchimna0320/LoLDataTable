"""Jungle Clear Compilation 社群刷野編纂表整合。

資料源為作者 Google 試算表的本機 xlsx 副本（人工放置、直接 hard code）：
- backend/data/jungle_clear_time/Jungle Clear Compilation S16_2026.xlsx
  （第一分頁 Jungle clear S162026，每列一種刷野）；
- 讀取前檢查檔案存在性，解析後快取為 JSON 供前端使用；
- 英雄 Q/W/E 技能圖示檔名取自 Data Dragon 英雄 JSON（依版本落盤快取）；
- 技能圖示本身由前端 /dd-img/spell/ 代理按需下載（與召喚師技能同路由）；
- 解析失敗時回退最近一次 JSON 快取，全失敗回傳 None。
"""
from __future__ import annotations

import datetime as _dt
import json
import re
from datetime import datetime
from functools import lru_cache

import openpyxl

from backend import config, data_access, ddragon_data
from backend.pipeline.common import (atomic_write_text, ensure_file_exists,
                                     get_logger, http_get, throttle)

logger = get_logger("jungle_data")

# 試算表第 4 列（index 3）才是表頭
_HEADER_ROW = 3
_COLUMNS = ["champion", "patch", "camps", "smites", "time", "path",
            "link", "player", "notes"]
_TIME_RE = re.compile(r"^(\d{1,2}):([0-5]\d)$")
_SKILL_RE = re.compile(r"\b[QWE]{3,4}\b")
_SHARD_CODES = "AD|AP|AS|AH|MS|HPF|HPS|T&SR"
_SHARD_RE = re.compile(rf"\b(?:{_SHARD_CODES})(?:/(?:{_SHARD_CODES})){{1,2}}\b")

# 營地英文→中文（配合站台中文版用語）
_CAMP_ZH = {
    "Red": "紅B", "Blue": "藍B", "Raptors": "六鳥", "Raptor": "六鳥",
    "Krugs": "石像", "Wolves": "三狼", "Gromp": "狼苔",
    "Scuttle": "河蟹", "Recall": "回城",
}
_RED_SIDE = {"Red", "Raptors", "Raptor", "Krugs"}
_BLUE_SIDE = {"Blue", "Wolves", "Gromp"}

# 符文碎片縮寫（Appendix 官方說明）
SHARD_ZH = {
    "AD": "攻擊力（適性）", "AP": "法強（適性）", "AS": "攻擊速度",
    "AH": "技能急速", "MS": "跑速", "HPF": "固定生命 +65",
    "HPS": "成長生命 +10~180", "T&SR": "韌性與減速抗性",
}


# ---------------------------------------------------------------------------
# 解析
# ---------------------------------------------------------------------------
def _to_seconds(text: str) -> int:
    match = _TIME_RE.match(text.strip())
    return int(match.group(1)) * 60 + int(match.group(2)) if match else 9999


def _parse_notes(notes: str) -> tuple[str, list[str], str]:
    """抽取起手技能順序、碎片代碼，其餘作為備註。"""
    skills = ""
    skill_match = _SKILL_RE.search(notes)
    if skill_match:
        skills = skill_match.group(0)
    shard_match = _SHARD_RE.search(notes)
    shards = shard_match.group(0).split("/") if shard_match else []
    remark = notes
    if skill_match:
        remark = remark.replace(skill_match.group(0), "")
    if shard_match:
        remark = remark.replace(shard_match.group(0), "")
    remark = re.sub(r"[\s.,;/]+", " ", remark).strip(" .;,/")
    return skills, shards, remark


def _translate_path(raw_path: str) -> tuple[str, str]:
    """營地路徑英→中，並判斷開局方（紅/藍/其他）。"""
    tokens = [t.strip() for t in re.split(r"\s*->\s*", raw_path) if t.strip()]
    zh = "→".join(_CAMP_ZH.get(t, t) for t in tokens)
    if tokens and tokens[0] in _RED_SIDE:
        side = "red"
    elif tokens and tokens[0] in _BLUE_SIDE:
        side = "blue"
    else:
        side = "other"
    return zh, side


def _cell_text(value) -> str:
    """xlsx 儲存格 → 去空白字串。"""
    return "" if value is None else str(value).strip()


def _normalize_patch(value) -> str:
    """版本欄：26.02→'26.02'、26.1→'26.1'；字串（如 26.01PBE）原樣保留。"""
    if isinstance(value, float):
        return f"{value:g}"
    return _cell_text(value)


def _int_text(value) -> str:
    """整數欄：xlsx 給 float（6.0）時轉成 '6'，空值回 ''。"""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(int(value))
    return _cell_text(value)


def _normalize_time(value) -> str:
    """時間欄：xlsx 為 datetime.time（分:秒），統一輸出 M:SS 字串。"""
    if isinstance(value, _dt.time):
        return f"{value.hour}:{value.minute:02d}"
    return _cell_text(value)


def _read_xlsx_rows(path) -> list[list[str]]:
    """讀本機 xlsx 第一分頁（S16），表頭後各列正規化為 9 個字串。

    Link 欄沒有文字時改取 Excel 超連結目標作備援。
    """
    # 檔案很小（約 380 列），用一般模式載入才能讀儲存格超連結
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = []
    for row in sheet.iter_rows(min_row=_HEADER_ROW + 2, max_col=len(_COLUMNS)):
        cells = list(row)
        link_cell = cells[6]
        link = _cell_text(link_cell.value)
        hyperlink = getattr(link_cell, "hyperlink", None)
        if not link and hyperlink is not None:
            link = _cell_text(hyperlink.target)
        rows.append([
            _cell_text(cells[0].value),          # champion
            _normalize_patch(cells[1].value),    # patch
            _int_text(cells[2].value),           # camps（稍後轉整數）
            _int_text(cells[3].value),           # smites
            _normalize_time(cells[4].value),     # time
            _cell_text(cells[5].value),          # path
            link,                                # link
            _cell_text(cells[7].value),          # player
            _cell_text(cells[8].value),          # notes
        ])
    workbook.close()
    return rows


def _parse_rows(rows: list[list[str]], valid_names: set[str]) -> list[dict]:
    """正規化列 → 刷野記錄列表（過濾區段標題與附錄列）。"""
    records, section, outdated = [], "", False
    for row in rows:
        name, patch, camps, smites, time_text, path, link, player, notes = row
        if name.endswith(":") and not time_text:
            section = name
            outdated = name.startswith("Outdated")
            continue
        if name not in valid_names or not _TIME_RE.match(time_text):
            continue
        skills, shards, remark = _parse_notes(notes)
        path_zh, side = _translate_path(path)
        records.append({
            "champion": name,
            "time": time_text,
            "seconds": _to_seconds(time_text),
            "patch": patch,
            "camps": int(camps) if camps.isdigit() else None,
            "smites": int(smites) if smites.isdigit() else None,
            "skills": list(skills),
            "shards": shards,
            "path_zh": path_zh,
            "side": side,
            "player": player,
            "link": link,
            "remark": remark,
            "outdated": outdated,
        })
    return records


# ---------------------------------------------------------------------------
# 技能圖示對應
# ---------------------------------------------------------------------------
def _load_name_map() -> dict[str, str]:
    """讀取英雄頭像同步時建立的「顯示名→DDragon ID」映射。"""
    if not config.DD_META_PATH.exists():
        return {}
    try:
        meta = json.loads(config.DD_META_PATH.read_text(encoding="utf-8"))
        return meta.get("name_to_id", {})
    except json.JSONDecodeError:
        return {}


def _champ_skill_icons(champ_id: str, version: str,
                       last_call: float) -> tuple[dict, float]:
    """取單一英雄 Q/W/E 圖示檔名；英雄 JSON 依版本落盤快取。"""
    target = config.DD_STATIC_DIR / "champ_json" / f"{champ_id}.{version}.json"
    if not target.exists():
        last_call = throttle(last_call)
        resp = http_get(config.DDRAGON_CHAMP_PAGE_URL.format(
            version=version, champ_id=champ_id))
        atomic_write_text(target, resp.text)
    try:
        detail = json.loads(target.read_text(encoding="utf-8"))
        spells = detail["data"][champ_id]["spells"][:3]
        icons = {key: spell.get("image", {}).get("full", "")
                 for key, spell in zip(("Q", "W", "E"), spells)}
    except (json.JSONDecodeError, KeyError) as exc:
        logger.warning("英雄 %s 技能資料解析失敗：%s", champ_id, exc)
        icons = {}
    return icons, last_call


def _aggregate(records: list[dict], version: str) -> list[dict]:
    """以英雄分組、補上技能圖示，回傳卡片排序後的列表。"""
    name_map = _load_name_map()
    grouped: dict[str, list[dict]] = {}
    for rec in records:
        grouped.setdefault(rec["champion"], []).append(rec)
    champions, last_call = [], 0.0
    for name, clears in grouped.items():
        clears.sort(key=lambda r: r["seconds"])
        champ_id = name_map.get(name, "")
        skill_icons = {}
        if champ_id:
            try:
                skill_icons, last_call = _champ_skill_icons(
                    champ_id, version, last_call)
            except Exception as exc:  # noqa: BLE001
                logger.warning("略過 %s 技能圖示：%s", name, exc)
        champions.append({
            "name": name,
            "champ_id": champ_id,
            "clears_n": len(clears),
            "fastest": clears[0]["time"],
            "skill_icons": skill_icons,
            "clears": clears,
        })
    champions.sort(key=lambda c: (-c["clears_n"],
                                  c["clears"][0]["seconds"], c["name"]))
    return champions


# ---------------------------------------------------------------------------
# 對外入口
# ---------------------------------------------------------------------------
def _valid_champion_names() -> set[str]:
    """有效英雄名：DDragon 顯示名＋人工覆蓋＋倉儲出現過的名稱。"""
    valid = set(_load_name_map()) | set(config.CHAMPION_NAME_OVERRIDES)
    try:
        valid |= set(data_access.champion_wall(None)["champion"])
    except Exception as exc:  # noqa: BLE001
        logger.warning("無法讀取倉儲英雄名，僅用 DDragon 映射：%s", exc)
    return valid


def build_cache() -> dict:
    """讀本機 xlsx 副本、解析並補技能圖示後寫入 JSON 快取。"""
    ensure_file_exists(config.JUNGLE_XLSX_PATH)
    rows = _read_xlsx_rows(config.JUNGLE_XLSX_PATH)
    records = _parse_rows(rows, _valid_champion_names())
    if not records:
        raise RuntimeError("刷野 xlsx 解析結果為 0 列，停止覆寫快取")
    version = ddragon_data.synced_version()
    champions = _aggregate(records, version)
    source_mtime = datetime.fromtimestamp(
        config.JUNGLE_XLSX_PATH.stat().st_mtime)
    payload = {
        "fetched_at": source_mtime.isoformat(timespec="seconds"),
        "source_file": config.JUNGLE_XLSX_PATH.name,
        "sheet_url": config.JUNGLE_SHEET_PAGE_URL,
        "discord_url": config.JUNGLE_DISCORD_URL,
        "version": version,
        "summary": {"champions": len(champions),
                    "clears": sum(c["clears_n"] for c in champions)},
        "champions": champions,
    }
    atomic_write_text(config.JUNGLE_JSON,
                      json.dumps(payload, ensure_ascii=False))
    logger.info("刷野編纂表已由本機副本更新：%d 英雄、%d 種刷野",
                len(champions), payload["summary"]["clears"])
    return payload


def _cached_payload() -> dict | None:
    """讀取磁碟上的 JSON 快取。"""
    if not config.JUNGLE_JSON.exists():
        return None
    try:
        return json.loads(config.JUNGLE_JSON.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        logger.error("刷野 JSON 快取毀損")
        return None


@lru_cache(maxsize=1)
def load_clears() -> dict | None:
    """載入刷野資料（每個伺服器進程僅計算一次）。

    xlsx 未變更時直接複用 JSON 快取（毫秒級），變更才重新解析。
    """
    cached = _cached_payload()
    try:
        mtime = datetime.fromtimestamp(
            config.JUNGLE_XLSX_PATH.stat().st_mtime)
        if cached and cached.get("fetched_at") == mtime.isoformat(
                timespec="seconds"):
            return cached
        return build_cache()
    except Exception as exc:  # noqa: BLE001
        logger.warning("本機刷野副本解析失敗（嘗試 JSON 快取）：%s", exc)
        if cached:
            return cached
        logger.error("刷野 JSON 快取不存在且本機副本無法讀取")
        return None
